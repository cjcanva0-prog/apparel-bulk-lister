import base64
import datetime
import io
import json
import os
import re
import zipfile
import rarfile
import openpyxl
import pandas as pd
from PIL import Image
from rapidfuzz import fuzz, process
import requests
import streamlit as st
from google import genai
from google.genai import types

DB_FILE = "catalog_db.json"

st.set_page_config(page_title="Apparel Multi-Marketplace Hub V2 (AI-Powered)", layout="wide")


# ==============================================================================
# GITHUB PERSISTENT SYNC (SECRETS-ENABLED)
# ==============================================================================
def get_github_headers():
    token = st.secrets.get("GITHUB_TOKEN", None)
    if token:
        return {
            "Authorization": f"token {token}",
            "Accept": "application/vnd.github.v3+json",
        }
    return None


def normalize_db_records(data):
    for k, v in data.items():
        if "color_variants" not in v:
            v["color_variants"] = [
                {
                    "color_name": v.get("color", "White"),
                    "color_map": v.get("color", "White"),
                    "images": v.get("images", []),
                }
            ]
        if "top_length" not in v:
            v["top_length"] = "Knee Length"
        if "amazon_overrides" not in v:
            v["amazon_overrides"] = {}
        if "myntra_overrides" not in v:
            v["myntra_overrides"] = {}

        # FIX: purge garbage values saved earlier by the INDIRECT-formula dropdown bug
        for ovr_key in ("amazon_overrides", "myntra_overrides"):
            v[ovr_key] = {
                k2: val
                for k2, val in v.get(ovr_key, {}).items()
                if not str(val).strip().upper().startswith("INDIRECT(")
            }
    return data


def load_db():
    headers = get_github_headers()
    if headers and "REPO_NAME" in st.secrets:
        repo = st.secrets["REPO_NAME"]
        fpath = st.secrets.get("FILE_PATH", "catalog_db.json")
        branch = st.secrets.get("BRANCH", "main")
        url = f"https://api.github.com/repos/{repo}/contents/{fpath}?ref={branch}"
        try:
            r = requests.get(url, headers=headers, timeout=10)
            if r.status_code == 200:
                content_b64 = r.json().get("content", "")
                data_str = base64.b64decode(content_b64).decode("utf-8")
                return normalize_db_records(json.loads(data_str))
        except Exception:
            pass

    if os.path.exists(DB_FILE):
        with open(DB_FILE, "r") as f:
            try:
                return normalize_db_records(json.load(f))
            except Exception:
                return {}
    return {}


def save_db(data):
    with open(DB_FILE, "w") as f:
        json.dump(data, f, indent=2)

    headers = get_github_headers()
    if headers and "REPO_NAME" in st.secrets:
        repo = st.secrets["REPO_NAME"]
        fpath = st.secrets.get("FILE_PATH", "catalog_db.json")
        branch = st.secrets.get("BRANCH", "main")
        url = f"https://api.github.com/repos/{repo}/contents/{fpath}"
        try:
            r_get = requests.get(url + f"?ref={branch}", headers=headers, timeout=10)
            sha = r_get.json().get("sha", None) if r_get.status_code == 200 else None

            content_str = json.dumps(data, indent=2)
            content_b64 = base64.b64encode(content_str.encode("utf-8")).decode("utf-8")

            payload = {
                "message": "Auto-sync catalog update from Streamlit Hub V2",
                "content": content_b64,
                "branch": branch,
            }
            if sha:
                payload["sha"] = sha

            r_put = requests.put(url, headers=headers, json=payload, timeout=10)
            if r_put.status_code in [200, 201]:
                st.toast("✅ Saved permanently to GitHub!", icon="🎉")
            else:
                st.error(f"GitHub Sync Error: {r_put.status_code} - {r_put.text}")
        except Exception as e:
            st.error(f"GitHub Sync Connection Failed: {e}")


db = load_db()

# ==============================================================================
# BASE MAPPERS & FALLBACK DROPDOWNS
# ==============================================================================
COLOR_MAPPER = {
    "White": "White", "Off White": "Off White", "Cream": "Off White",
    "Beige": "Beige", "Black": "Black", "Navy Blue": "Blue", "Blue": "Blue",
    "Teal": "Turquoise", "Turquoise Blue": "Turquoise", "Green": "Green",
    "Olive": "Green", "Mehendi Green": "Green", "Sea Green": "Green",
    "Lime Green": "Green", "Mustard": "Yellow", "Yellow": "Yellow",
    "Maroon": "Red", "Red": "Red", "Burgundy": "Red", "Wine": "Red",
    "Rust": "Orange", "Orange": "Orange", "Coral": "Pink", "Peach": "Pink",
    "Pink": "Pink", "Rani Pink": "Pink", "Fuchsia": "Pink", "Magenta": "Pink",
    "Purple": "Purple", "Lavender": "Purple", "Mauve": "Purple",
    "Brown": "Brown", "Coffee Brown": "Brown", "Charcoal": "Grey",
    "Grey": "Grey", "Gold": "Gold", "Silver": "Silver", "Bronze": "Bronze",
    "Copper": "Bronze", "Multi": "Multicolor", "Other / Custom...": "Other / Custom...",
}

FABRIC_MAPPER = {
    "Pure Cotton": "Cotton", "Cotton Blend": "Cotton Blend",
    "Cotton Silk": "Silk Blend", "Silk Blend": "Silk Blend",
    "Georgette": "Georgette", "Chanderi": "Chanderi", "Rayon": "Rayon",
    "Organza": "Organza", "Tissue": "Tissue", "Modal": "Modal",
    "Satin": "Satin", "Linen Blend": "Linen Blend",
}

LENGTH_MAPPER = {
    "Knee Length": ("Knee Length", "Knee Length"),
    "Calf Length": ("Calf Length", "Calf Length"),
    "Above Knee": ("Above Knee", "Mid Thigh Length"),
    "Floor Length": ("Floor Length", "Floor Length"),
    "Ankle Length": ("Calf Length", "Ankle Length"),
    "Short / Hip Length": ("Short", "Hip Length"),
    "Other / Custom...": ("Knee Length", "Knee Length"),
}

AMAZON_SIZE_MAP = {
    "XS": "XS (x_s)", "S": "S (s)", "M": "M (m)", "L": "L (l)",
    "XL": "XL", "2XL": "2XL (2x_l)", "XXL": "2XL (2x_l)",
    "3XL": "3XL", "4XL": "4XL", "5XL": "5XL"
}

AMAZON_EXACT_DESCRIPTION = (
    "Elevate your ethnic with this elegant Women's Kurta , designed to blend comfort with timeless style. "
    "Made from a soft and breathable  fabric, this Kurta featuring delicate white floral embroidery on the sleeves and neckline, "
    "for a refined look. The neck design, 3/4 elasticated sleeves, provide effortless comfort while maintaining a flattering look. "
    "Whether you're dressing for work, casual outings, festive occasions, family gatherings, or travel, this kurta keeps you looking "
    "polished throughout the day. Style it with statement earrings, bangles, and sandals for a graceful ethnic, or keep it minimal for a chic everyday look."
)

DROPDOWNS = {
    "colors": list(COLOR_MAPPER.keys()),
    "amazon_color_maps": [
        "Beige", "Black", "Blue", "Bronze", "Brown", "Clear", "Gold",
        "Green", "Grey", "Metallic", "Multicolor", "Off White", "Orange",
        "Pink", "Purple", "Red", "Silver", "Turquoise", "White", "Yellow",
    ],
    "top_lengths": list(LENGTH_MAPPER.keys()),
    "fabrics": list(FABRIC_MAPPER.keys()) + ["Other / Custom..."],
    "top_patterns": [
        "Embroidered", "Printed", "Solid", "Woven Design", "Yoke Design",
        "Self Design", "Striped", "Checked", "Colourblocked", "Other / Custom...",
    ],
    "print_types": [
        "Floral", "Geometric", "Paisley", "Ethnic Motifs", "Abstract", "Solid",
        "Tribal", "Chevron", "Tie and Dye", "Polka Dot", "Other / Custom...",
    ],
    "neck_styles": [
        "V-Neck", "Round Neck", "Mandarin Collar", "Sweetheart Neck", "Boat Neck",
        "Square Neck", "Halter Neck", "Shirt Collar", "Scoop Neck", "Keyhole Neck",
        "Other / Custom...",
    ],
    "sleeve_lengths": [
        "Three-Quarter Sleeves", "Short Sleeves", "Long Sleeves", "Sleeveless",
        "Other / Custom...",
    ],
    "shapes": [
        "Straight", "A-Line", "Anarkali", "Flared", "Kaftan", "Pathani",
        "Other / Custom...",
    ],
    "weave_types": [
        "Machine Weave", "Handloom", "Regular", "Knitted", "Powerloom",
        "Other / Custom...",
    ],
    "wash_cares": [
        "Dry Clean", "Hand Wash", "Machine Wash", "Dry Clean Only",
        "Hand Wash Only", "Other / Custom...",
    ],
    "occasions": [
        "Festive", "Casual", "Daily", "Party", "Wedding", "Fusion", "Work",
        "Other / Custom...",
    ],
    "packages": [
        "1 Kurta, 1 Pant", "1 Kurta, 1 Pant, 1 Dupatta", "1 Kurta",
        "1 Kurta, 1 Salwar", "1 Kurta, 1 Palazzo", "1 Kurta, 1 Sharara, 1 Dupatta",
        "Other / Custom...",
    ],
}


# ==============================================================================
# V2 ENGINE: DYNAMIC TEMPLATE INTROSPECTION & DROPDOWN EXTRACTOR
# ==============================================================================
def _cells_from_range(ws, rng):
    """Returns non-empty cell values from a range. Handles single cells like $F$4."""
    rng = rng.replace("$", "")
    got = ws[rng]
    if not isinstance(got, tuple):  # single cell -> openpyxl returns a Cell, not a tuple
        got = ((got,),)
    out = []
    for row in got:
        for c in (row if isinstance(row, tuple) else (row,)):
            if c.value is not None and str(c.value).strip() != "":
                out.append(str(c.value).strip())
    return out


def _values_from_reference(wb, ref):
    """Resolves a defined name or 'Sheet'!$A$1:$A$9 reference into a list of values."""
    ref = str(ref).strip()
    if ref.startswith("="):
        ref = ref[1:]
    if ref in wb.defined_names:  # named range -> real range
        ref = wb.defined_names[ref].attr_text
    if "!" not in ref:
        return []
    sheet_part, range_part = ref.rsplit("!", 1)
    sheet_name = sheet_part.strip("'").replace("''", "'")
    if sheet_name not in wb.sheetnames:
        return []
    try:
        vals = _cells_from_range(wb[sheet_name], range_part)
    except Exception:
        return []
    return list(dict.fromkeys(vals))


def resolve_excel_range(wb, formula_str, product_type="KURTA"):
    """
    Turns a data-validation formula into a list of dropdown options.
    Supports:
      1) inline lists wrapped in quotes:   "A,B,C"
      2) Amazon dependent dropdowns:       INDIRECT(<product type cell> & "attr_name...")
         (resolved by substituting the product type, e.g. KURTA + "stylemarketplace_id...")
      3) named ranges and Sheet!A1:A9 ranges
    """
    if not formula_str:
        return []
    f = str(formula_str).strip()
    if f.startswith("="):
        f = f[1:]

    # 1) Inline literal list — ONLY when wrapped in quotes
    if f.startswith('"') and f.endswith('"'):
        return [o.strip() for o in f[1:-1].split(",") if o.strip()]

    # 2) Amazon INDIRECT(...) dependent dropdown
    if f.upper().startswith("INDIRECT("):
        m = re.search(r'&\s*"([^"]+)"\s*\)\s*$', f)
        if not m:
            return []
        pt = (product_type or "").replace("-", "_").replace(" ", "")
        if pt[:1].isdigit():
            pt = "_" + pt
        return _values_from_reference(wb, pt + m.group(1))

    # 3) Named range / plain sheet range
    return _values_from_reference(wb, f)


def parse_marketplace_template(uploaded_file, platform="Amazon"):
    uploaded_file.seek(0)
    wb = openpyxl.load_workbook(uploaded_file, data_only=True)

    if platform == "Myntra":
        ws = wb["Kurta Sets"] if "Kurta Sets" in wb.sheetnames else (wb["Kurtas"] if "Kurtas" in wb.sheetnames else wb.active)
        header_row = 3
        product_type = ""
    else:
        ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
        header_row = 4
        product_type = "KURTA"

    schema = {}
    col_to_header = {}

    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val:
            h_name = str(val).strip()
            if h_name not in schema:
                schema[h_name] = {"col_idx": col, "options": []}
                col_to_header[col] = h_name

    if ws.data_validations and ws.data_validations.dataValidation:
        for dv in ws.data_validations.dataValidation:
            if dv.type == "list" and dv.formula1:
                opts = resolve_excel_range(wb, dv.formula1, product_type)
                if opts:
                    for cell_range in dv.sqref.ranges:
                        for c_idx in range(cell_range.min_col, cell_range.max_col + 1):
                            if c_idx in col_to_header:
                                h_name = col_to_header[c_idx]
                                if not schema[h_name]["options"]:
                                    schema[h_name]["options"] = opts

    for sname in wb.sheetnames:
        if any(w in sname.lower() for w in ["valid", "drop", "master", "value"]):
            v_ws = wb[sname]
            for col in range(1, min(v_ws.max_column + 1, 150)):
                v_head = v_ws.cell(row=1, column=col).value or v_ws.cell(row=2, column=col).value
                if v_head:
                    v_head_clean = str(v_head).strip()
                    matched_key = None
                    if v_head_clean in schema:
                        matched_key = v_head_clean
                    elif schema:
                        m = process.extractOne(v_head_clean, list(schema.keys()), scorer=fuzz.WRatio)
                        if m and m[1] >= 90:
                            matched_key = m[0]
                    if matched_key and not schema[matched_key]["options"]:
                        opts = []
                        for r in range(2, min(v_ws.max_row + 1, 500)):
                            cval = v_ws.cell(row=r, column=col).value
                            if cval is not None and str(cval).strip() and str(cval).strip() != v_head_clean:
                                opts.append(str(cval).strip())
                        if opts:
                            schema[matched_key]["options"] = list(dict.fromkeys(opts))

    uploaded_file.seek(0)
    return schema


def snap_to_valid_option(target_val: str, options: list) -> str:
    """Snaps any string value to the exact or closest matching option in `options`.
    Returns "" (blank) when nothing is a reasonable match — never a random first option."""
    if target_val is None or str(target_val).strip() == "":
        return ""
    target_str = str(target_val).strip()
    if not options:
        return target_str

    if target_str in options:
        return target_str

    lower_map = {o.lower(): o for o in options}
    if target_str.lower() in lower_map:
        return lower_map[target_str.lower()]

    best = process.extractOne(target_str, options, scorer=fuzz.WRatio)
    if best and best[1] >= 60:
        return best[0]
    return ""


def match_value_to_list(val: str, options: list, default_val: str = "") -> str:
    if not val or not options:
        return default_val
    clean_opts = [o for o in options if o != "Other / Custom..."]
    snapped = snap_to_valid_option(val, clean_opts)
    return snapped if snapped else default_val


# ==============================================================================
# ONE-WAY REACTIVE CASCADE: CORE ATTRIBUTES -> AMAZON & MYNTRA TABS
# ==============================================================================
def compute_baseline_template_maps(core_dict):
    """
    Computes the full baseline dictionary for BOTH Amazon and Myntra templates
    from the current Core Attributes dictionary.
    """
    d_code = core_dict.get("design_code", "NewStyle")
    t_core = core_dict.get("title_core", "Floral Cotton Embroidered Kurta Set with Pants")
    raw_fab = core_dict.get("fabric", "Pure Cotton")
    amz_fab = FABRIC_MAPPER.get(raw_fab, raw_fab)
    pat = core_dict.get("top_pattern", "Embroidered")
    prt = core_dict.get("print_type", "Floral")
    nck = core_dict.get("neck", "V-Neck")
    slv = core_dict.get("sleeve_length", "Three-Quarter Sleeves")
    amz_slv = "3/4 Sleeve" if "Three" in str(slv) else str(slv)
    shp = core_dict.get("shape", "Straight")
    user_top_len = core_dict.get("top_length", "Calf Length")
    len_tuple = LENGTH_MAPPER.get(user_top_len, (user_top_len, user_top_len))
    myn_len, amz_len = len_tuple[0], len_tuple[1]
    wev = core_dict.get("weave", "Machine Weave")
    wsh = core_dict.get("wash_care", "Dry Clean")
    occ = core_dict.get("occasion", "Festive")
    mrp = str(core_dict.get("mrp", 3999))
    sp = str(core_dict.get("selling_price", 1499))
    hsn = str(core_dict.get("hsn", "62114210"))
    pkg = core_dict.get("package_contains", "1 Kurta, 1 Pant")
    desc = core_dict.get("description", AMAZON_EXACT_DESCRIPTION)

    brand = "PERVAS"
    brand_email = "pervas0007@gmail.com"
    title_full = f"{brand} Women's {raw_fab} {pat} {t_core} ({d_code})"

    amz_base = {
        "Status": "Active",
        "Title": title_full,
        "SKU": f"{brand}-{d_code}-Parent",
        "Product Type": "KURTA",
        "Listing Action": "Create or Replace (Full Update)",
        "Parentage Level": "Child",
        "Parent SKU": f"{brand}-{d_code}-Parent",
        "Variation Theme Name": "SIZE/COLOR",
        "Item Name": title_full,
        "Brand Name": brand,
        "Product Id Type": "GTIN Exempt",
        "Model Name": str(d_code),
        "Part Number": str(d_code),
        "Manufacturer": brand_email,
        "Manufacturer Contact Information": brand_email,
        "Packer Contact Information": brand_email,
        "Product Description": desc,
        "Generic Keywords": "two piece suit set coord dress stylish co ord sets cord coords dresses a line neck ords long salwar trendy new printed latest design traditional suits ladies daily rayon festive kutis indian kutties v kurthi fashion mul weddings floral print",
        "Lifestyle": "Casual" if occ in ["Casual", "Daily", "Work"] else "Ethnic",
        "Style": shp,
        "Fit Type": "Regular",
        "Department Name": "Womens",
        "Target Gender": "Female",
        "Age Range Description": "Adult",
        "Apparel Size System": "IN",
        "Apparel Size Class": "Alpha",
        "Shirt Size System": "IN",
        "Shirt Size Class": "Alpha",
        "Shirt Body Type": "Regular",
        "Special Size": "Standard",
        "Fabric Type": amz_fab,
        "Material": amz_fab,
        "Pattern": prt,
        "Design Name": pat,
        "Item Type Name": "KURTA",
        "Item Length Description": amz_len,
        "Occasion": occ,
        "Care Instructions": wsh,
        "Collar Style": nck,
        "Neck Style": nck,
        "Sleeve Length Description": amz_slv,
        "Weave Method": "Powerloom" if wev == "Machine Weave" else wev,
        "External Product Information Entity": "HSN Code",
        "External Product Information": hsn,
        "Standard Price": sp,
        "Maximum Retail Price": mrp,
        "Your Price INR (Sell on Amazon, IN)": sp,
        "Maximum Retail Price (Sell on Amazon, IN)": mrp,
        "Item Weight": "450.0",
        "Item Weight Unit": "Grams",
        "Country of Origin": "India",
        "Skip Offer": "No",
        "Item Condition": "New",
        "Offer Condition Note": "New",
        "Fulfillment Channel Code (IN)": "AMAZON_IN",
        "Number of Items": "1",
        "Item Length Longer Edge": "46.0",
        "Item Length Unit": "Centimeters",
        "Package Length": "25.0",
        "Package Length Unit": "Centimeters",
        "Package Width": "22.0",
        "Package Width Unit": "Centimeters",
        "Package Height": "3.0",
        "Package Height Unit": "Centimeters",
        "Package Weight": "450.0",
        "Package Weight Unit": "Grams",
    }

    myn_base = {
        "brand": brand,
        "vendorArticleName": f"{brand} {t_core}",
        "productDisplayName": f"{brand} {t_core}",
        "Manufacturer Name and Address with Pincode": "Pervas, 2087, The Palladium Mall, Yogi Chowk, Surat, Gujarat - 395010",
        "Packer Name and Address with Pincode": "Pervas, 2087, The Palladium Mall, Yogi Chowk, Surat, Gujarat - 395010",
        "Country Of Origin": "India",
        "articleType": "Kurta Sets" if "Pant" in pkg or "Sharara" in pkg or "Palazzo" in pkg or "Salwar" in pkg else "Kurtas",
        "is Standard Size present on Label": "Yes",
        "HSN": hsn,
        "MRP": mrp,
        "ISP": sp,
        "AgeGroup": "Adults-Women",
        "FashionType": "Fashion",
        "Usage": "Ethnic" if occ in ["Festive", "Wedding", "Party"] else "Casual",
        "Year": str(datetime.datetime.now().year),
        "Product Details": desc,
        "materialCareDescription": wsh,
        "Top Fabric": raw_fab,
        "Top Pattern": pat,
        "Neck": nck,
        "Sleeve Length": slv,
        "Top Shape": shp,
        "Top Length": myn_len,
        "Bottom Fabric": raw_fab,
        "Bottom Pattern": pat,
        "Bottom Closure": "Elasticated",
        "Waistband": "Elasticated",
        "Print or Pattern Type": prt,
        "Occasion": occ,
        "Weave Pattern": "Regular",
        "Weave Type": wev,
        "Wash Care": wsh,
        "Stitch": "Ready to Wear",
        "Add-Ons": "NA",
        "Number of Pockets": "NA",
        "Number of Items": "1",
        "Net Quantity Unit": "Piece",
        "Package Contains": pkg,
        "Net Quantity": "1",
    }
    return amz_base, myn_base


def set_template_widget_value(platform_prefix, field_name, raw_value, schema_dict):
    """
    Sets the Streamlit session_state key for an Amazon or Myntra field,
    automatically snapping to a valid dropdown option if that field has dropdown options.
    """
    w_key = f"{platform_prefix}_gran_{field_name}"
    meta = schema_dict.get(field_name, {})
    opts = meta.get("options", [])
    if opts:
        snapped = snap_to_valid_option(raw_value, opts)
        st.session_state[w_key] = snapped if snapped in opts else ""
    else:
        st.session_state[w_key] = "" if raw_value is None else str(raw_value)


def get_current_core_values():
    """Reads the current live values of Tab 1 (Core Attributes) from st.session_state."""
    def resolve_custom(sel_key, cust_key, default_v):
        sel = st.session_state.get(sel_key, default_v)
        return st.session_state.get(cust_key, default_v) if sel == "Other / Custom..." else sel

    return {
        "design_code": st.session_state.get("core_design_code", "NewStyle"),
        "title_core": st.session_state.get("core_title_core", "Floral Cotton Embroidered Kurta Set with Pants"),
        "fabric": resolve_custom("core_fabric_sel", "core_fabric_cust", "Pure Cotton"),
        "top_pattern": resolve_custom("core_pattern_sel", "core_pattern_cust", "Embroidered"),
        "print_type": resolve_custom("core_print_sel", "core_print_cust", "Floral"),
        "neck": resolve_custom("core_neck_sel", "core_neck_cust", "V-Neck"),
        "sleeve_length": resolve_custom("core_sleeve_sel", "core_sleeve_cust", "Three-Quarter Sleeves"),
        "shape": resolve_custom("core_shape_sel", "core_shape_cust", "Straight"),
        "top_length": resolve_custom("core_length_sel", "core_length_cust", "Calf Length"),
        "weave": resolve_custom("core_weave_sel", "core_weave_cust", "Machine Weave"),
        "wash_care": resolve_custom("core_wash_sel", "core_wash_cust", "Dry Clean"),
        "occasion": resolve_custom("core_occ_sel", "core_occ_cust", "Festive"),
        "mrp": st.session_state.get("core_mrp", 3999),
        "selling_price": st.session_state.get("core_sp", 1499),
        "hsn": st.session_state.get("core_hsn", "62114210"),
        "package_contains": resolve_custom("core_pkg_sel", "core_pkg_cust", "1 Kurta, 1 Pant"),
        "sizes": st.session_state.get("core_sizes", ["M", "L", "XL", "2XL", "3XL"]),
        "description": st.session_state.get("core_desc", AMAZON_EXACT_DESCRIPTION),
    }


def sync_all_core_to_templates(amz_schema, myn_schema, ai_amz_ovr=None, ai_myn_ovr=None):
    """
    Pushes ALL Core Attributes + Baseline Rules + AI Overrides into the
    Amazon and Myntra template widget states.
    """
    core_now = get_current_core_values()
    amz_base, myn_base = compute_baseline_template_maps(core_now)

    if ai_amz_ovr:
        for k, v in ai_amz_ovr.items():
            if v is not None and str(v).strip() != "":
                amz_base[k] = v

    if ai_myn_ovr:
        for k, v in ai_myn_ovr.items():
            if v is not None and str(v).strip() != "":
                myn_base[k] = v

    for fname in amz_schema.keys():
        val_to_set = amz_base.get(fname, None)
        if val_to_set is None:
            m = process.extractOne(fname, list(amz_base.keys()), scorer=fuzz.token_sort_ratio)
            if m and m[1] >= 88:
                val_to_set = amz_base[m[0]]
        if val_to_set is not None:
            set_template_widget_value("amz", fname, val_to_set, amz_schema)

    for fname in myn_schema.keys():
        val_to_set = myn_base.get(fname, None)
        if val_to_set is None:
            m = process.extractOne(fname, list(myn_base.keys()), scorer=fuzz.token_sort_ratio)
            if m and m[1] >= 88:
                val_to_set = myn_base[m[0]]
        if val_to_set is not None:
            set_template_widget_value("myn", fname, val_to_set, myn_schema)


def on_core_field_change(core_field_name):
    """
    One-Way Reactive Callback: Triggered whenever the user edits ANY single field in Tab 1 (Core Attributes).
    Immediately updates the dependent Amazon and Myntra fields in Tabs 2 & 3,
    without touching unrelated manual edits in Tabs 2 & 3!
    """
    amz_schema = st.session_state.get("amz_schema", {})
    myn_schema = st.session_state.get("myn_schema", {})
    core_now = get_current_core_values()
    amz_base, myn_base = compute_baseline_template_maps(core_now)

    CORE_TO_AMZ_MAP = {
        "design_code": ["SKU", "Parent SKU", "Model Name", "Part Number", "Title", "Item Name"],
        "title_core": ["Title", "Item Name"],
        "fabric": ["Fabric Type", "Material", "Title", "Item Name"],
        "top_pattern": ["Design Name", "Title", "Item Name"],
        "print_type": ["Pattern"],
        "neck": ["Neck Style", "Collar Style"],
        "sleeve_length": ["Sleeve Length Description"],
        "shape": ["Style"],
        "top_length": ["Item Length Description"],
        "weave": ["Weave Method"],
        "wash_care": ["Care Instructions"],
        "occasion": ["Occasion", "Lifestyle"],
        "mrp": ["Maximum Retail Price", "Maximum Retail Price (Sell on Amazon, IN)"],
        "selling_price": ["Standard Price", "Your Price INR (Sell on Amazon, IN)"],
        "hsn": ["External Product Information"],
        "description": ["Product Description"],
    }

    CORE_TO_MYN_MAP = {
        "title_core": ["vendorArticleName", "productDisplayName"],
        "fabric": ["Top Fabric", "Bottom Fabric"],
        "top_pattern": ["Top Pattern", "Bottom Pattern"],
        "print_type": ["Print or Pattern Type"],
        "neck": ["Neck"],
        "sleeve_length": ["Sleeve Length"],
        "shape": ["Top Shape"],
        "top_length": ["Top Length"],
        "weave": ["Weave Type"],
        "wash_care": ["Wash Care", "materialCareDescription"],
        "occasion": ["Occasion", "Usage"],
        "mrp": ["MRP"],
        "selling_price": ["ISP"],
        "hsn": ["HSN"],
        "package_contains": ["Package Contains", "articleType"],
        "description": ["Product Details"],
    }

    for target_amz_col in CORE_TO_AMZ_MAP.get(core_field_name, []):
        if target_amz_col in amz_base:
            for s_col in amz_schema.keys():
                if s_col == target_amz_col or fuzz.token_sort_ratio(s_col, target_amz_col) >= 90:
                    set_template_widget_value("amz", s_col, amz_base[target_amz_col], amz_schema)

    for target_myn_col in CORE_TO_MYN_MAP.get(core_field_name, []):
        if target_myn_col in myn_base:
            for s_col in myn_schema.keys():
                if s_col == target_myn_col or fuzz.token_sort_ratio(s_col, target_myn_col) >= 90:
                    set_template_widget_value("myn", s_col, myn_base[target_myn_col], myn_schema)


def load_record_into_session_state(record, amz_schema, myn_schema):
    """Initializes all Tab 1, Tab 2, and Tab 3 widget keys in st.session_state."""
    st.session_state["core_design_code"] = record.get("design_code", "NewStyle")
    st.session_state["core_title_core"] = record.get("title_core", "Floral Cotton Embroidered Kurta Set with Pants")
    st.session_state["core_fabric_sel"] = match_value_to_list(record.get("fabric", "Pure Cotton"), DROPDOWNS["fabrics"], "Pure Cotton")
    st.session_state["core_pattern_sel"] = match_value_to_list(record.get("top_pattern", "Embroidered"), DROPDOWNS["top_patterns"], "Embroidered")
    st.session_state["core_print_sel"] = match_value_to_list(record.get("print_type", "Floral"), DROPDOWNS["print_types"], "Floral")
    st.session_state["core_neck_sel"] = match_value_to_list(record.get("neck", "V-Neck"), DROPDOWNS["neck_styles"], "V-Neck")
    st.session_state["core_sleeve_sel"] = match_value_to_list(record.get("sleeve_length", "Three-Quarter Sleeves"), DROPDOWNS["sleeve_lengths"], "Three-Quarter Sleeves")
    st.session_state["core_shape_sel"] = match_value_to_list(record.get("shape", "Straight"), DROPDOWNS["shapes"], "Straight")
    st.session_state["core_length_sel"] = match_value_to_list(record.get("top_length", "Calf Length"), DROPDOWNS["top_lengths"], "Calf Length")
    st.session_state["core_weave_sel"] = match_value_to_list(record.get("weave", "Machine Weave"), DROPDOWNS["weave_types"], "Machine Weave")
    st.session_state["core_wash_sel"] = match_value_to_list(record.get("wash_care", "Dry Clean"), DROPDOWNS["wash_cares"], "Dry Clean")
    st.session_state["core_occ_sel"] = match_value_to_list(record.get("occasion", "Festive"), DROPDOWNS["occasions"], "Festive")
    st.session_state["core_mrp"] = int(record.get("mrp", 3999))
    st.session_state["core_sp"] = int(record.get("selling_price", 1499))
    st.session_state["core_hsn"] = str(record.get("hsn", "62114210"))
    st.session_state["core_pkg_sel"] = match_value_to_list(record.get("package_contains", "1 Kurta, 1 Pant"), DROPDOWNS["packages"], "1 Kurta, 1 Pant")
    valid_sz = ["XS", "S", "M", "L", "XL", "2XL", "3XL"]
    st.session_state["core_sizes"] = [s for s in record.get("sizes", ["M", "L", "XL", "2XL", "3XL"]) if s in valid_sz] or ["M", "L", "XL", "2XL"]
    st.session_state["core_desc"] = record.get("description", AMAZON_EXACT_DESCRIPTION)

    sync_all_core_to_templates(
        amz_schema,
        myn_schema,
        ai_amz_ovr=record.get("amazon_overrides", {}),
        ai_myn_ovr=record.get("myntra_overrides", {}),
    )


# ==============================================================================
# V2 ENGINE: ZIP/RAR ARCHIVE EXTRACTOR & GEMINI AI MULTIMODAL ANALYZER
# ==============================================================================
def extract_archive(uploaded_archive):
    texts = []
    images = []
    raw_bytes = io.BytesIO(uploaded_archive.read())
    is_rar = uploaded_archive.name.lower().endswith(".rar")
    archive_cls = rarfile.RarFile if is_rar else zipfile.ZipFile

    with archive_cls(raw_bytes) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            fname = info.filename
            lower_name = fname.lower()
            if "__macosx" in lower_name or os.path.basename(fname).startswith("."):
                continue

            if lower_name.endswith((".txt", ".md", ".csv", ".json")):
                content = z.read(fname).decode("utf-8", errors="ignore")
                texts.append(f"[{os.path.basename(fname)}]:\n{content}")
            elif lower_name.endswith((".jpg", ".jpeg", ".png", ".webp")):
                img_data = z.read(fname)
                img = Image.open(io.BytesIO(img_data)).convert("RGB")
                images.append((os.path.basename(fname), img))

    uploaded_archive.seek(0)
    return "\n\n".join(texts), images


def run_gemini_archive_analysis(api_key, model_name, raw_text, images, amz_schema, myn_schema, archive_name=""):
    client = genai.Client(api_key=api_key)

    compressed_imgs = []
    for _, img in images[:6]:
        im_copy = img.copy()
        im_copy.thumbnail((800, 800))
        compressed_imgs.append(im_copy)

    amz_fields_spec = {
        k: (v["options"][:60] if v["options"] else "Free Text")
        for k, v in amz_schema.items()
    }
    myn_fields_spec = {
        k: (v["options"][:60] if v["options"] else "Free Text")
        for k, v in myn_schema.items()
    }

    prompt = f"""
    You are an expert Indian ethnic wear catalog manager for Amazon.in and Myntra.
    Analyze the attached product images and the raw text notes from the archive '{archive_name}'.

    RAW TEXT NOTES FROM ARCHIVE:
    {raw_text if raw_text else "No text file found. Infer all details from the archive name and product images."}

    BASE CATALOG DROPDOWN OPTIONS:
    - fabrics: {DROPDOWNS['fabrics'][:-1]}
    - top_patterns: {DROPDOWNS['top_patterns'][:-1]}
    - print_types: {DROPDOWNS['print_types'][:-1]}
    - neck_styles: {DROPDOWNS['neck_styles'][:-1]}
    - sleeve_lengths: {DROPDOWNS['sleeve_lengths'][:-1]}
    - shapes: {DROPDOWNS['shapes'][:-1]}
    - top_lengths: {DROPDOWNS['top_lengths'][:-1]}
    - weave_types: {DROPDOWNS['weave_types'][:-1]}
    - wash_cares: {DROPDOWNS['wash_cares'][:-1]}
    - occasions: {DROPDOWNS['occasions'][:-1]}
    - packages: {DROPDOWNS['packages'][:-1]}
    - colors: {DROPDOWNS['colors'][:-1]}
    - amazon_color_maps: {DROPDOWNS['amazon_color_maps']}

    AMAZON TEMPLATE FIELDS & ALLOWED DROPDOWNS:
    {json.dumps(amz_fields_spec) if amz_fields_spec else "No custom Amazon template loaded yet."}

    MYNTRA TEMPLATE FIELDS & ALLOWED DROPDOWNS:
    {json.dumps(myn_fields_spec) if myn_fields_spec else "No custom Myntra template loaded yet."}

    CRITICAL INSTRUCTIONS:
    1. Extract 'design_code' from SKU/Design name in text or archive filename (capitalize nicely, e.g., 'Bansari').
    2. Identify ALL distinct color variants shown across the uploaded images (e.g., if 5 photos show 5 different colors, list all 5 in 'detected_colors').
    3. Parse size range (e.g., 'M to XXXL' -> ["M", "L", "XL", "2XL", "3XL"]). Note: If 'Price' in notes is wholesale/cost price (like 350), set 'selling_price' to a realistic retail price or that number, and 'mrp' to 3999.
    4. PREFILL ALL RELEVANT AMAZON & MYNTRA TEMPLATE FIELDS: In 'amazon_overrides' and 'myntra_overrides', populate every relevant attribute from the template schemas above (including Collar Style, Sleeve Type, Fit Type, Closure Type, Bottom Style, Ornamentation, Weave Type, Occasion, Care Instructions, Lifestyle, Item Highlight, Bullet Points, etc.).
    5. NEXT-BEST VALUE RULE: Whenever a field has allowed dropdown options, you MUST pick the exact match or the NEXT BEST, most semantically relevant option from its allowed options list.

    Return ONLY valid JSON matching this structure:
    {{
      "base_fields": {{
        "design_code": "string",
        "title_core": "string",
        "fabric": "string",
        "top_pattern": "string",
        "print_type": "string",
        "neck": "string",
        "sleeve_length": "string",
        "shape": "string",
        "top_length": "string",
        "weave": "string",
        "wash_care": "string",
        "occasion": "string",
        "mrp": 3999,
        "selling_price": 1499,
        "hsn": "62114210",
        "package_contains": "string",
        "sizes": ["M", "L", "XL", "2XL", "3XL"],
        "description": "string"
      }},
      "detected_colors": [
        {{"color_name": "Navy Blue", "color_map": "Blue"}}
      ],
      "amazon_overrides": {{
        "Exact Amazon Column Name": "Prefilled Value"
      }},
      "myntra_overrides": {{
        "Exact Myntra Column Name": "Prefilled Value"
      }}
    }}
    """

    response = client.models.generate_content(
        model=model_name,
        contents=[prompt] + compressed_imgs,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            temperature=0.1,
        ),
    )
    return json.loads(response.text)


# ==============================================================================
# SIDEBAR: AI STUDIO KEY & GLOBAL TEMPLATE SCHEMA LOADER
# ==============================================================================
with st.sidebar:
    st.header("🤖 AI & Template Engine (V2)")

    default_key = st.secrets.get("GEMINI_API_KEY", "") if "GEMINI_API_KEY" in st.secrets else ""
    gemini_api_key = st.text_input(
        "Gemini API Key (Google AI Studio)",
        value=default_key,
        type="password",
        help="Get a 100% free key from aistudio.google.com or store GEMINI_API_KEY in Streamlit Secrets.",
    )

    model_choice = st.selectbox(
        "Free-Tier Gemini Model",
        [
            "gemini-3.8-flash",
            "gemini-3.5-flash-lite",
            "gemini-3.5-flash",
            "Custom Model ID...",
        ],
        index=0,
        help="Free tier in Google AI Studio — multimodal vision + structured JSON.",
    )
    if model_choice == "Custom Model ID...":
        gemini_model = st.text_input("Enter Custom Model ID", value="gemini-3.8-flash")
    else:
        gemini_model = model_choice

    st.markdown("---")
    st.subheader("📑 Load Master Templates")
    st.caption("Upload your Amazon & Myntra templates here so the Add/Edit form can display and prefill every field and dropdown from both platforms.")

    tpl_amz_sidebar = st.file_uploader("Amazon Master Template (.xlsm/.xlsx)", type=["xlsm", "xlsx"], key="sb_amz_tpl")
    tpl_myn_sidebar = st.file_uploader("Myntra Master Template (.xlsx)", type=["xlsx"], key="sb_myn_tpl")

    if tpl_amz_sidebar:
        if st.session_state.get("loaded_amz_tpl_name") != tpl_amz_sidebar.name:
            with st.spinner("Extracting Amazon columns & dropdowns..."):
                st.session_state["amz_schema"] = parse_marketplace_template(tpl_amz_sidebar, "Amazon")
                st.session_state["loaded_amz_tpl_name"] = tpl_amz_sidebar.name
                if st.session_state.get("show_add_modal") or st.session_state.get("edit_product_key"):
                    ek = st.session_state.get("edit_product_key")
                    saved_amz = db.get(ek, {}).get("amazon_overrides", {}) if ek and ek in db else {}
                    sync_all_core_to_templates(st.session_state["amz_schema"], st.session_state.get("myn_schema", {}), ai_amz_ovr=saved_amz)
        st.success(f"Amazon: {len(st.session_state.get('amz_schema', {}))} fields loaded")

    if tpl_myn_sidebar:
        if st.session_state.get("loaded_myn_tpl_name") != tpl_myn_sidebar.name:
            with st.spinner("Extracting Myntra columns & dropdowns..."):
                st.session_state["myn_schema"] = parse_marketplace_template(tpl_myn_sidebar, "Myntra")
                st.session_state["loaded_myn_tpl_name"] = tpl_myn_sidebar.name
                if st.session_state.get("show_add_modal") or st.session_state.get("edit_product_key"):
                    ek = st.session_state.get("edit_product_key")
                    saved_myn = db.get(ek, {}).get("myntra_overrides", {}) if ek and ek in db else {}
                    sync_all_core_to_templates(st.session_state.get("amz_schema", {}), st.session_state["myn_schema"], ai_myn_ovr=saved_myn)
        st.success(f"Myntra: {len(st.session_state.get('myn_schema', {}))} fields loaded")

amz_schema = st.session_state.get("amz_schema", {})
myn_schema = st.session_state.get("myn_schema", {})

st.title("👗 Apparel Catalog & Multi-Marketplace Hub (V2 AI)")

# ==============================================================================
# 1. CATALOG MANAGEMENT (ADD / EDIT / ZIP AI INGESTION)
# ==============================================================================
col_t1, col_t2, col_t3, col_t4 = st.columns([2, 1, 1, 1])
with col_t1:
    st.subheader("📦 Master Catalog")
with col_t2:
    if st.button("➕ Add New Product (AI / Manual)", use_container_width=True, type="primary"):
        st.session_state["show_add_modal"] = True
        st.session_state["edit_product_key"] = None
        st.session_state["temp_colors"] = [{"color_name": "White", "color_map": "White", "images": ""}]
        load_record_into_session_state({}, amz_schema, myn_schema)
        st.rerun()
with col_t3:
    if st.button("🗑️ Delete Selected", use_container_width=True):
        st.session_state["trigger_delete"] = True
with col_t4:
    catalog_json_str = json.dumps(db, indent=2)
    st.download_button(
        label="💾 Backup DB",
        data=catalog_json_str,
        file_name="catalog_db.json",
        mime="application/json",
        use_container_width=True,
    )

# Add / Edit Drawer (Reactive — No st.form freeze!)
if st.session_state.get("show_add_modal", False) or st.session_state.get("edit_product_key", None):
    edit_key = st.session_state.get("edit_product_key", None)
    is_edit = edit_key is not None and edit_key in db

    if is_edit and st.session_state.get("loaded_edit_key") != edit_key:
        curr_rec = db[edit_key]
        st.session_state["temp_colors"] = [
            {
                "color_name": c.get("color_name", "White"),
                "color_map": c.get("color_map", "White"),
                "images": "\n".join(c.get("images", [])),
            }
            for c in curr_rec.get("color_variants", [])
        ]
        for idx_c, c_var in enumerate(st.session_state["temp_colors"]):
            st.session_state[f"cname_{idx_c}"] = match_value_to_list(c_var["color_name"], DROPDOWNS["colors"], "White")
            st.session_state[f"cmap_{idx_c}"] = match_value_to_list(c_var["color_map"], DROPDOWNS["amazon_color_maps"], "White")
            st.session_state[f"cimg_{idx_c}"] = c_var["images"]

        load_record_into_session_state(curr_rec, amz_schema, myn_schema)
        st.session_state["loaded_edit_key"] = edit_key

    st.info(f"✏️ **{'Editing Style: ' + edit_key if is_edit else 'Add New Design to Catalog (Live One-Way Core -> Template Sync)'}**")

    # --------------------------------------------------------------------------
    # AI ARCHIVE DROPZONE (.ZIP / .RAR)
    # --------------------------------------------------------------------------
    with st.expander("✨ **Auto-Prefill from Product Archive (.zip / .rar)** — Drop folder like `Love Sharara.rar` or `Bansari.zip`", expanded=not is_edit):
        up_col1, up_col2 = st.columns([3, 1])
        with up_col1:
            product_archive = st.file_uploader(
                "Upload Product Folder (.zip or .rar containing images & text notes)",
                type=["zip", "rar"],
                key="archive_uploader",
            )
        with up_col2:
            st.write("")
            st.write("")
            run_ai_btn = st.button("🪄 Analyze & Prefill with AI", type="primary", use_container_width=True)

        if run_ai_btn:
            if not product_archive:
                st.warning("Please upload a .zip or .rar file first.")
            elif not gemini_api_key:
                st.error("Please enter your free Gemini API Key in the left sidebar.")
            else:
                with st.spinner("Extracting archive & analyzing product photos + notes with Gemini..."):
                    try:
                        raw_txt, extracted_imgs = extract_archive(product_archive)
                        st.session_state["last_extracted_imgs"] = extracted_imgs
                        st.session_state["last_extracted_txt"] = raw_txt

                        ai_res = run_gemini_archive_analysis(
                            api_key=gemini_api_key,
                            model_name=gemini_model,
                            raw_text=raw_txt,
                            images=extracted_imgs,
                            amz_schema=amz_schema,
                            myn_schema=myn_schema,
                            archive_name=product_archive.name,
                        )

                        bf = ai_res.get("base_fields", {})
                        bf["amazon_overrides"] = ai_res.get("amazon_overrides", {})
                        bf["myntra_overrides"] = ai_res.get("myntra_overrides", {})

                        # Load directly into widget session_state keys (Core + All Amazon + All Myntra!)
                        load_record_into_session_state(bf, amz_schema, myn_schema)

                        det_colors = ai_res.get("detected_colors", [])
                        if det_colors:
                            new_temp_colors = []
                            for idx_c, dc in enumerate(det_colors):
                                cname = match_value_to_list(dc.get("color_name"), DROPDOWNS["colors"], "White")
                                cmap = match_value_to_list(dc.get("color_map"), DROPDOWNS["amazon_color_maps"], COLOR_MAPPER.get(cname, "White"))
                                new_temp_colors.append({"color_name": cname, "color_map": cmap, "images": ""})
                                st.session_state[f"cname_{idx_c}"] = cname
                                st.session_state[f"cmap_{idx_c}"] = cmap
                                st.session_state[f"cimg_{idx_c}"] = ""
                            st.session_state["temp_colors"] = new_temp_colors

                        st.toast("✅ AI Prefilled Core Attributes + Amazon & Myntra Template Tabs!", icon="🪄")
                        st.rerun()
                    except Exception as e:
                        st.error(f"AI Analysis Error: {e}")

        if st.session_state.get("last_extracted_imgs") or st.session_state.get("last_extracted_txt"):
            st.caption("📸 **Extracted Archive Preview:**")
            if st.session_state.get("last_extracted_txt"):
                st.code(st.session_state["last_extracted_txt"], language="text")
            imgs_preview = st.session_state.get("last_extracted_imgs", [])
            if imgs_preview:
                p_cols = st.columns(min(len(imgs_preview), 6))
                for idx, (iname, iobj) in enumerate(imgs_preview[:6]):
                    p_cols[idx].image(iobj, caption=iname, use_container_width=True)

    # --------------------------------------------------------------------------
    # LIVE REACTIVE TABS (ONE-WAY CASCADE: TAB 1 -> TAB 2 & TAB 3)
    # --------------------------------------------------------------------------
    tab_core, tab_amz_all, tab_myn_all = st.tabs([
        "🏷️ 1. Core Attributes & Colorways (Auto-Updates Amazon & Myntra)",
        f"📦 2. All Amazon Template Fields ({len(amz_schema)})",
        f"👗 3. All Myntra Template Fields ({len(myn_schema)})",
    ])

    with tab_core:
        f1, f2, f3 = st.columns(3)
        with f1:
            st.markdown("### 🏷️ Base Identifiers")
            st.text_input("Design Code / Model Name*", key="core_design_code", on_change=on_core_field_change, args=("design_code",))
            st.text_input("Core Title (Without Brand)*", key="core_title_core", on_change=on_core_field_change, args=("title_core",))

            st.selectbox("Fabric*", DROPDOWNS["fabrics"], key="core_fabric_sel", on_change=on_core_field_change, args=("fabric",))
            if st.session_state.get("core_fabric_sel") == "Other / Custom...":
                st.text_input("Type Custom Fabric", key="core_fabric_cust", on_change=on_core_field_change, args=("fabric",))

            st.selectbox("Pattern / Work*", DROPDOWNS["top_patterns"], key="core_pattern_sel", on_change=on_core_field_change, args=("top_pattern",))
            if st.session_state.get("core_pattern_sel") == "Other / Custom...":
                st.text_input("Type Custom Pattern", key="core_pattern_cust", on_change=on_core_field_change, args=("top_pattern",))

            st.selectbox("Print Type*", DROPDOWNS["print_types"], key="core_print_sel", on_change=on_core_field_change, args=("print_type",))
            if st.session_state.get("core_print_sel") == "Other / Custom...":
                st.text_input("Type Custom Print", key="core_print_cust", on_change=on_core_field_change, args=("print_type",))

        with f2:
            st.markdown("### ✂️ Cut & Styling")
            st.selectbox("Neck Type*", DROPDOWNS["neck_styles"], key="core_neck_sel", on_change=on_core_field_change, args=("neck",))
            if st.session_state.get("core_neck_sel") == "Other / Custom...":
                st.text_input("Type Custom Neck", key="core_neck_cust", on_change=on_core_field_change, args=("neck",))

            st.selectbox("Sleeve Length*", DROPDOWNS["sleeve_lengths"], key="core_sleeve_sel", on_change=on_core_field_change, args=("sleeve_length",))
            if st.session_state.get("core_sleeve_sel") == "Other / Custom...":
                st.text_input("Type Custom Sleeve", key="core_sleeve_cust", on_change=on_core_field_change, args=("sleeve_length",))

            st.selectbox("Kurta Shape / Fit*", DROPDOWNS["shapes"], key="core_shape_sel", on_change=on_core_field_change, args=("shape",))
            if st.session_state.get("core_shape_sel") == "Other / Custom...":
                st.text_input("Type Custom Shape", key="core_shape_cust", on_change=on_core_field_change, args=("shape",))

            st.selectbox("Kurta / Top Length*", DROPDOWNS["top_lengths"], key="core_length_sel", on_change=on_core_field_change, args=("top_length",))
            if st.session_state.get("core_length_sel") == "Other / Custom...":
                st.text_input("Type Custom Length", key="core_length_cust", on_change=on_core_field_change, args=("top_length",))

            st.selectbox("Weave Type*", DROPDOWNS["weave_types"], key="core_weave_sel", on_change=on_core_field_change, args=("weave",))
            if st.session_state.get("core_weave_sel") == "Other / Custom...":
                st.text_input("Type Custom Weave", key="core_weave_cust", on_change=on_core_field_change, args=("weave",))

            st.selectbox("Wash Care*", DROPDOWNS["wash_cares"], key="core_wash_sel", on_change=on_core_field_change, args=("wash_care",))
            if st.session_state.get("core_wash_sel") == "Other / Custom...":
                st.text_input("Type Custom Wash Care", key="core_wash_cust", on_change=on_core_field_change, args=("wash_care",))

            st.selectbox("Occasion*", DROPDOWNS["occasions"], key="core_occ_sel", on_change=on_core_field_change, args=("occasion",))
            if st.session_state.get("core_occ_sel") == "Other / Custom...":
                st.text_input("Type Custom Occasion", key="core_occ_cust", on_change=on_core_field_change, args=("occasion",))

        with f3:
            st.markdown("### 💰 Pricing & Standard Curve")
            st.number_input("MRP (₹)*", step=100, key="core_mrp", on_change=on_core_field_change, args=("mrp",))
            st.number_input("Selling Price (₹)*", step=50, key="core_sp", on_change=on_core_field_change, args=("selling_price",))
            st.text_input("HSN Code*", key="core_hsn", on_change=on_core_field_change, args=("hsn",))

            st.selectbox("Package Contains*", DROPDOWNS["packages"], key="core_pkg_sel", on_change=on_core_field_change, args=("package_contains",))
            if st.session_state.get("core_pkg_sel") == "Other / Custom...":
                st.text_input("Type Custom Package", key="core_pkg_cust", on_change=on_core_field_change, args=("package_contains",))

            st.multiselect(
                "Available Sizes*",
                ["XS", "S", "M", "L", "XL", "2XL", "3XL"],
                key="core_sizes",
            )

        st.markdown("---")
        st.markdown("### 🎨 Color Variants & Image Sets")
        updated_color_variants = []
        for idx, cv in enumerate(st.session_state.get("temp_colors", [])):
            st.markdown(f"**Colorway #{idx + 1}**")
            c_col1, c_col2, c_col3 = st.columns([1.5, 1.5, 3])
            with c_col1:
                if f"cname_{idx}" not in st.session_state:
                    st.session_state[f"cname_{idx}"] = match_value_to_list(cv.get("color_name", "White"), DROPDOWNS["colors"], "White")
                col_name_sel = st.selectbox(f"Color Name #{idx + 1}", DROPDOWNS["colors"], key=f"cname_{idx}")
                if col_name_sel == "Other / Custom...":
                    col_name = st.text_input(f"Type Custom Color #{idx + 1}", key=f"ccustom_{idx}")
                else:
                    col_name = col_name_sel

            with c_col2:
                suggested_map = COLOR_MAPPER.get(col_name_sel, cv.get("color_map", "Multicolor"))
                if f"cmap_{idx}" not in st.session_state:
                    st.session_state[f"cmap_{idx}"] = match_value_to_list(suggested_map, DROPDOWNS["amazon_color_maps"], "White")
                col_map_choice = st.selectbox(
                    f"Amazon Filter Bucket #{idx + 1}",
                    DROPDOWNS["amazon_color_maps"],
                    key=f"cmap_{idx}",
                )

            with c_col3:
                if f"cimg_{idx}" not in st.session_state:
                    st.session_state[f"cimg_{idx}"] = cv.get("images", "")
                img_text = st.text_area(f"Images for {col_name} (1 URL per line)", key=f"cimg_{idx}", height=70)

            updated_color_variants.append(
                {
                    "color_name": col_name,
                    "color_map": col_map_choice,
                    "images": [u.strip() for u in str(img_text).splitlines() if u.strip()],
                }
            )

        col_btn1, col_btn2 = st.columns([1, 4])
        with col_btn1:
            if st.button("➕ Add Another Color Variant"):
                new_i = len(st.session_state["temp_colors"])
                st.session_state["temp_colors"].append({"color_name": "White", "color_map": "White", "images": ""})
                st.session_state[f"cname_{new_i}"] = "White"
                st.session_state[f"cmap_{new_i}"] = "White"
                st.session_state[f"cimg_{new_i}"] = ""
                st.rerun()
        with col_btn2:
            if len(st.session_state.get("temp_colors", [])) > 1 and st.button("➖ Remove Last Color Variant"):
                st.session_state["temp_colors"].pop()
                st.rerun()

        st.markdown("---")
        st.text_area(
            "Product Description",
            key="core_desc",
            height=120,
            on_change=on_core_field_change,
            args=("description",),
        )

    # --------------------------------------------------------------------------
    # TAB 2: ALL AMAZON TEMPLATE FIELDS (INDEPENDENT ONE-WAY TARGET)
    # --------------------------------------------------------------------------
    amz_override_inputs = {}
    with tab_amz_all:
        if not amz_schema:
            st.info("👈 Upload your **Amazon Master Template (.xlsm / .xlsx)** in the left sidebar to unlock granular control over every Amazon column and dropdown.")
        else:
            st.caption("Every column from your Amazon template is prefilled below from AI + Core Attributes. Editing Core Attributes updates these automatically, while edits you make here stay strictly in Amazon!")
            amz_grid = st.columns(3)
            for idx, (fname, fmeta) in enumerate(amz_schema.items()):
                col_ui = amz_grid[idx % 3]
                w_key = f"amz_gran_{fname}"
                opts = fmeta.get("options", [])
                if opts:
                    full_opts = [""] + opts
                    if w_key not in st.session_state or st.session_state[w_key] not in full_opts:
                        st.session_state[w_key] = ""
                    amz_override_inputs[fname] = col_ui.selectbox(
                        f"📦 {fname}",
                        options=full_opts,
                        key=w_key,
                    )
                else:
                    if w_key not in st.session_state:
                        st.session_state[w_key] = ""
                    amz_override_inputs[fname] = col_ui.text_input(
                        f"📦 {fname}",
                        key=w_key,
                    )

    # --------------------------------------------------------------------------
    # TAB 3: ALL MYNTRA TEMPLATE FIELDS (INDEPENDENT ONE-WAY TARGET)
    # --------------------------------------------------------------------------
    myn_override_inputs = {}
    with tab_myn_all:
        if not myn_schema:
            st.info("👈 Upload your **Myntra Master Template (.xlsx)** in the left sidebar to unlock granular control over every Myntra column and dropdown.")
        else:
            st.caption("Every column from your Myntra template is prefilled below from AI + Core Attributes. Editing Core Attributes updates these automatically, while edits you make here stay strictly in Myntra!")
            myn_grid = st.columns(3)
            for idx, (fname, fmeta) in enumerate(myn_schema.items()):
                col_ui = myn_grid[idx % 3]
                w_key = f"myn_gran_{fname}"
                opts = fmeta.get("options", [])
                if opts:
                    full_opts = [""] + opts
                    if w_key not in st.session_state or st.session_state[w_key] not in full_opts:
                        st.session_state[w_key] = ""
                    myn_override_inputs[fname] = col_ui.selectbox(
                        f"👗 {fname}",
                        options=full_opts,
                        key=w_key,
                    )
                else:
                    if w_key not in st.session_state:
                        st.session_state[w_key] = ""
                    myn_override_inputs[fname] = col_ui.text_input(
                        f"👗 {fname}",
                        key=w_key,
                    )

    st.markdown("---")
    save_c1, save_c2 = st.columns([1, 4])
    with save_c1:
        save_btn = st.button("💾 Save Product to Catalog", type="primary", use_container_width=True)
    with save_c2:
        cancel_btn = st.button("Cancel", use_container_width=False)

    if cancel_btn:
        st.session_state["show_add_modal"] = False
        st.session_state["edit_product_key"] = None
        st.session_state["loaded_edit_key"] = None
        st.rerun()

    if save_btn:
        core_final = get_current_core_values()
        d_code = core_final["design_code"]
        sz_list = core_final["sizes"]

        std_measurements = {
            "XS": {"Across Shoulder": 13.5, "Bust": 34.0, "Chest": 34.0, "Front Length": 29.0, "Hips": 32.0, "Waist": 28.0, "Inseam Length": 25.0},
            "S": {"Across Shoulder": 14.0, "Bust": 36.0, "Chest": 36.0, "Front Length": 29.0, "Hips": 34.0, "Waist": 30.0, "Inseam Length": 25.0},
            "M": {"Across Shoulder": 14.5, "Bust": 38.0, "Chest": 38.0, "Front Length": 29.0, "Hips": 36.0, "Waist": 32.0, "Inseam Length": 24.8},
            "L": {"Across Shoulder": 15.0, "Bust": 40.0, "Chest": 40.0, "Front Length": 29.0, "Hips": 38.0, "Waist": 34.0, "Inseam Length": 24.3},
            "XL": {"Across Shoulder": 15.5, "Bust": 42.0, "Chest": 42.0, "Front Length": 29.0, "Hips": 40.0, "Waist": 36.0, "Inseam Length": 23.8},
            "2XL": {"Across Shoulder": 16.0, "Bust": 44.0, "Chest": 44.0, "Front Length": 29.0, "Hips": 42.0, "Waist": 38.0, "Inseam Length": 22.0},
            "3XL": {"Across Shoulder": 16.5, "Bust": 46.0, "Chest": 46.0, "Front Length": 29.0, "Hips": 44.0, "Waist": 40.0, "Inseam Length": 22.0},
        }

        existing_amz_ovr = db.get(d_code, {}).get("amazon_overrides", {})
        existing_myn_ovr = db.get(d_code, {}).get("myntra_overrides", {})

        clean_amz_ovr = (
            {k: v for k, v in amz_override_inputs.items() if str(v).strip() != ""}
            if amz_schema else existing_amz_ovr
        )
        clean_myn_ovr = (
            {k: v for k, v in myn_override_inputs.items() if str(v).strip() != ""}
            if myn_schema else existing_myn_ovr
        )

        db[d_code] = {
            **core_final,
            "net_qty": 1,
            "measurements": {sz: std_measurements.get(sz, {}) for sz in sz_list},
            "color_variants": updated_color_variants,
            "amazon_overrides": clean_amz_ovr,
            "myntra_overrides": clean_myn_ovr,
        }
        save_db(db)
        st.session_state["show_add_modal"] = False
        st.session_state["edit_product_key"] = None
        st.session_state["loaded_edit_key"] = None
        st.rerun()

# ==============================================================================
# 2. CATALOG TABLE
# ==============================================================================
if not db:
    st.info("No styles found in catalog. Click '➕ Add New Product' above to create one.")
else:
    table_data = []
    for key, item in db.items():
        colors_list = [c.get("color_name") for c in item.get("color_variants", [])]
        table_data.append({
            "Select": False,
            "Design Code": item.get("design_code"),
            "Title": item.get("title_core"),
            "Colors": ", ".join(colors_list) if colors_list else item.get("color", ""),
            "Fabric": item.get("fabric"),
            "Pattern": item.get("top_pattern"),
            "Neck": item.get("neck"),
            "Sleeve": item.get("sleeve_length"),
            "Shape": item.get("shape"),
            "Top Length": item.get("top_length", "Knee Length"),
            "MRP (₹)": item.get("mrp"),
            "Price (₹)": item.get("selling_price"),
            "Sizes": ", ".join(item.get("sizes", [])),
        })

    df_catalog = pd.DataFrame(table_data)

    edited_df = st.data_editor(
        df_catalog,
        column_config={
            "Select": st.column_config.CheckboxColumn("Select", default=False),
            "Design Code": st.column_config.TextColumn("Design Code", disabled=True),
            "Title": st.column_config.TextColumn("Title", disabled=True),
        },
        disabled=[c for c in df_catalog.columns if c != "Select"],
        hide_index=True,
        use_container_width=True,
    )

    selected_designs = edited_df[edited_df["Select"] == True]["Design Code"].tolist()

    if st.session_state.get("trigger_delete", False):
        if not selected_designs:
            st.warning("Please tick the checkbox next to the product(s) you want to delete.")
        else:
            for d in selected_designs:
                if d in db:
                    del db[d]
            save_db(db)
            st.session_state["trigger_delete"] = False
            st.rerun()

    if len(selected_designs) == 1:
        if st.button(f"✏️ Edit Selected ({selected_designs[0]})"):
            st.session_state["edit_product_key"] = selected_designs[0]
            st.session_state["show_add_modal"] = False
            st.session_state["loaded_edit_key"] = None
            st.rerun()

    st.markdown("---")

    # ==============================================================================
    # 3. LISTING HUB (AUTOMATED PARENT-CHILD + GRANULAR OVERRIDE INJECTION)
    # ==============================================================================
    st.subheader("🚀 List Selected Products")

    if not selected_designs:
        st.write("👈 *Tick one or more products in the catalog table above to configure and generate listing files.*")
    else:
        st.success(f"Selected **{len(selected_designs)}** product(s): **{', '.join(selected_designs)}**")

        c_plat, c_brand, c_tpl = st.columns([1.5, 2, 2.5])

        with c_plat:
            marketplace = st.selectbox("Select Target Marketplace", ["Myntra", "Amazon.in"])
            default_cat = "Kurta Sets" if marketplace == "Myntra" else "KURTA"
            category_val = st.text_input("Category / Product Type", value=default_cat)

        with c_brand:
            if marketplace == "Myntra":
                selected_brands = st.multiselect(
                    "Select Brands to Generate",
                    ["KALINI", "MITERA", "PERVAS", "PERABIN"],
                    default=["KALINI", "MITERA", "PERVAS", "PERABIN"],
                )
            else:
                selected_brands = st.multiselect(
                    "Select Brand(s)",
                    ["PERVAS", "PERABIN"],
                    default=["PERVAS"],
                )

        with c_tpl:
            uploaded_template = st.file_uploader(
                f"Upload Latest {marketplace} Template (.xlsx / .xlsm)",
                type=["xlsx", "xlsm"],
            )

        if st.button(f"⚡ Generate {marketplace} Bulk Upload File for {len(selected_designs)} Product(s)", type="primary", use_container_width=True):
            active_tpl = uploaded_template or (tpl_myn_sidebar if marketplace == "Myntra" else tpl_amz_sidebar)

            if not active_tpl:
                st.error(f"Please upload the {marketplace} template file (here or in the sidebar).")
            elif not selected_brands:
                st.error("Please select at least one brand.")
            else:
                active_tpl.seek(0)
                is_xlsm = active_tpl.name.lower().endswith(".xlsm")
                keep_vba_flag = True if is_xlsm else False
                out_ext = ".xlsm" if is_xlsm else ".xlsx"
                out_mime = "application/vnd.ms-excel.sheet.macroEnabled.12" if is_xlsm else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

                wb = openpyxl.load_workbook(active_tpl, keep_vba=keep_vba_flag)

                if marketplace == "Myntra":
                    sheet_name = "Kurta Sets" if "set" in category_val.lower() else "Kurtas"
                    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
                    header_row = 3
                    start_row = 4
                else:
                    ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
                    header_row = 4
                    start_row = 7
                    for c in range(1, ws.max_column + 1):
                        ws.cell(row=6, column=c).value = None

                col_map = {}
                header_names = []
                for c in range(1, ws.max_column + 1):
                    h_val = ws.cell(row=header_row, column=c).value
                    if h_val:
                        col_name_clean = str(h_val).strip()
                        if col_name_clean not in col_map:
                            col_map[col_name_clean] = []
                            header_names.append(col_name_clean)
                        col_map[col_name_clean].append(c)

                def write_cell(row_idx, col_name, value, occurrence=0):
                    if value is None or str(value).strip() == "":
                        return
                    target_col_name = None
                    if col_name in col_map:
                        target_col_name = col_name
                    elif header_names:
                        match, score, _ = process.extractOne(col_name, header_names, scorer=fuzz.token_sort_ratio)
                        if score >= 88:
                            target_col_name = match
                    if target_col_name and target_col_name in col_map:
                        cols = col_map[target_col_name]
                        if occurrence < len(cols):
                            ws.cell(row=row_idx, column=cols[occurrence], value=value)

                current_row = start_row
                group_id_counter = 1

                for d_name in selected_designs:
                    prod = db[d_name]
                    colorways = prod.get("color_variants", [{"color_name": "White", "color_map": "White", "images": []}])
                    amz_ovr = prod.get("amazon_overrides", {})
                    myn_ovr = prod.get("myntra_overrides", {})
                    amz_base_defaults, myn_base_defaults = compute_baseline_template_maps(prod)

                    raw_fabric = prod.get("fabric", "Pure Cotton")

                    for brand in selected_brands:
                        brand_email = "pervas0007@gmail.com" if brand == "PERVAS" else "perabin32544@gmail.com"

                        # ------------------------------------------------------
                        # AMAZON PARENT ROW
                        # ------------------------------------------------------
                        if marketplace == "Amazon.in":
                            parent_sku = f"{brand}-{prod.get('design_code')}-Parent"
                            parent_title = f"{brand} Women's {raw_fabric} {prod.get('top_pattern', 'Printed')} Kurta Pant Set ({prod.get('design_code')})"

                            p_row = {
                                **amz_base_defaults,
                                **amz_ovr,
                                "Status": "Active",
                                "Title": parent_title,
                                "SKU": parent_sku,
                                "Product Type": category_val.upper(),
                                "Parentage Level": "Parent",
                                "Parent SKU": "",
                                "Variation Theme Name": "SIZE/COLOR",
                                "Item Name": parent_title,
                                "Brand Name": brand,
                                "Manufacturer": brand_email,
                                "Manufacturer Contact Information": brand_email,
                                "Packer Contact Information": brand_email,
                            }

                            for col_name, val in p_row.items():
                                write_cell(current_row, col_name, val)

                            first_c_imgs = colorways[0].get("images", [])
                            if first_c_imgs:
                                write_cell(current_row, "Main Image URL", first_c_imgs[0])
                                for i_idx, img_url in enumerate(first_c_imgs[1:8]):
                                    write_cell(current_row, "Other Image URL", img_url, occurrence=i_idx)

                            write_cell(current_row, "Bullet Point", "1. Elegant Design: Features sophisticated embroidery and styling for everyday and festive elegance.", occurrence=0)
                            write_cell(current_row, "Bullet Point", "2. Soft & Breathable Fabric: Crafted from lightweight fabric offering all-day comfort.", occurrence=1)
                            write_cell(current_row, "Bullet Point", "3. Versatile Occasion: Ideal for casual outings, office wear, family gatherings, and celebrations.", occurrence=2)
                            write_cell(current_row, "Bullet Point", "4. Easy Care & Lasting Style: Durable stitching and long-lasting fabric quality.", occurrence=3)

                            current_row += 1

                        # ------------------------------------------------------
                        # CHILD ROWS: LOOP THROUGH EVERY COLORWAY & SIZE
                        # ------------------------------------------------------
                        for cway in colorways:
                            c_name = cway.get("color_name", "White")
                            c_map = cway.get("color_map", c_name)
                            c_imgs = cway.get("images", [])

                            for sz in prod.get("sizes", []):
                                amz_size_val = AMAZON_SIZE_MAP.get(sz.upper(), sz)
                                myntra_size = "XXL" if sz.upper() in ["2XL", "XXL"] else sz
                                m = prod.get("measurements", {}).get(sz, {})

                                if marketplace == "Myntra":
                                    sku = f"{brand}-P-{prod.get('design_code')}-{c_name}-{myntra_size}"
                                    art_num = f"{brand}-P-{prod.get('design_code')}-{c_name}"
                                    display_name = f"{brand} {c_name} {prod.get('title_core', '')}"

                                    style_note_text = (
                                        f"Step into elegance with this mesmerizing Kurta set by {brand}. "
                                        f"Whether you are heading to a family dinner or a festive celebration, the lush {raw_fabric} fabric "
                                        f"creates a look that is both traditional and trendy. The {prod.get('neck', 'classic')} neckline adds a touch of modern grace, "
                                        f"while the flattering silhouette ensures you stay comfortable all day long. "
                                        f"Pair this ensemble with oxidized jewelry and block heels for a complete festive look."
                                    )

                                    m_row = {
                                        **myn_base_defaults,
                                        **myn_ovr,
                                        "styleGroupId": group_id_counter,
                                        "vendorSkuCode": sku,
                                        "vendorArticleNumber": art_num,
                                        "vendorArticleName": display_name,
                                        "productDisplayName": display_name,
                                        "brand": brand,
                                        "articleType": category_val,
                                        "Brand Size": myntra_size,
                                        "Standard Size": myntra_size,
                                        "Brand Colour (Remarks)": c_name,
                                        "Prominent Colour": c_map,
                                        "SKUCode": sku,
                                        "styleNote": myn_ovr.get("styleNote") or style_note_text,
                                        "Across Shoulder ( Inches )": m.get("Across Shoulder"),
                                        "Bust ( Inches )": m.get("Bust"),
                                        "Chest ( Inches )": m.get("Chest"),
                                        "Front Length ( Inches )": m.get("Front Length"),
                                        "Hips ( Inches )": m.get("Hips"),
                                        "Waist ( Inches )": m.get("Waist"),
                                        "Inseam Length ( Inches )": m.get("Inseam Length"),
                                    }

                                    for col_name, val in m_row.items():
                                        write_cell(current_row, col_name, val)

                                else:
                                    child_sku = f"{brand}-{prod.get('design_code')}-{c_name}-{sz}"
                                    parent_sku = f"{brand}-{prod.get('design_code')}-Parent"
                                    child_title = f"{brand} Women's {raw_fabric} Kurta Pant Set ({prod.get('design_code')} {c_name} {sz})"

                                    c_row = {
                                        **amz_base_defaults,
                                        **amz_ovr,
                                        "Status": "Active",
                                        "Title": child_title,
                                        "SKU": child_sku,
                                        "Product Type": category_val.upper(),
                                        "Parentage Level": "Child",
                                        "Parent SKU": parent_sku,
                                        "Variation Theme Name": "SIZE/COLOR",
                                        "Item Name": child_title,
                                        "Brand Name": brand,
                                        "Apparel Size Value": amz_size_val,
                                        "Shirt Size Value": amz_size_val,
                                        "Color": c_name,
                                        "Color Map": c_map,
                                        "Manufacturer": brand_email,
                                        "Manufacturer Contact Information": brand_email,
                                        "Packer Contact Information": brand_email,
                                    }

                                    for col_name, val in c_row.items():
                                        write_cell(current_row, col_name, val)

                                    if c_imgs:
                                        write_cell(current_row, "Main Image URL", c_imgs[0])
                                        for i_idx, img_url in enumerate(c_imgs[1:8]):
                                            write_cell(current_row, "Other Image URL", img_url, occurrence=i_idx)

                                    write_cell(current_row, "Bullet Point", "1. Elegant Design: Features sophisticated embroidery and styling for everyday and festive elegance.", occurrence=0)
                                    write_cell(current_row, "Bullet Point", "2. Soft & Breathable Fabric: Crafted from lightweight fabric offering all-day comfort.", occurrence=1)
                                    write_cell(current_row, "Bullet Point", "3. Versatile Occasion: Ideal for casual outings, office wear, family gatherings, and celebrations.", occurrence=2)
                                    write_cell(current_row, "Bullet Point", "4. Easy Care & Lasting Style: Durable stitching and long-lasting fabric quality.", occurrence=3)

                                current_row += 1
                            group_id_counter += 1

                output = io.BytesIO()
                wb.save(output)
                output.seek(0)

                total_rows = current_row - start_row
                st.session_state["generated_file"] = {
                    "data": output.getvalue(),
                    "filename": f"{marketplace}_bulk_listing_{len(selected_designs)}_styles{out_ext}",
                    "mime": out_mime,
                    "rows": total_rows,
                    "designs_count": len(selected_designs),
                    "brands_count": len(selected_brands),
                    "marketplace": marketplace,
                }

        if "generated_file" in st.session_state:
            gen = st.session_state["generated_file"]
            st.success(f"🎉 Generated {gen['rows']} total rows across {gen['designs_count']} style(s) and {gen['brands_count']} brand(s)!")
            st.download_button(
                label=f"📥 Download Bulk Upload Sheet ({gen['marketplace']})",
                data=gen["data"],
                file_name=gen["filename"],
                mime=gen["mime"],
                use_container_width=True,
            )
